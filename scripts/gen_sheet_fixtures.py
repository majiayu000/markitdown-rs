"""Write the generated spreadsheet fixtures in tests/fixtures/xlsx/gen-*.xlsx.

Most workbooks are written as raw SpreadsheetML so every cell type, style
and quirk is exact; one is written with openpyxl for a realistic package.

usage: uv run --project oracle python scripts/gen_sheet_fixtures.py
"""
import datetime as dt
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "tests/fixtures/xlsx"
NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def col_name(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def workbook(name, sheets, *, numfmts=(), xfs=(0,), shared=(), date1904=False):
    """sheets: [(name, rows)], rows: [(row_number or None, [cell xml or None, ...])]
    cell xml is the inside of <c> plus attributes: (attrs, inner)."""
    files = {}
    overrides = ""
    sheet_entries = ""
    wb_rels = ""
    for i, (sname, rows) in enumerate(sheets, 1):
        body = []
        for rnum, cells in rows:
            r_attr = f' r="{rnum}"' if rnum is not None else ""
            cs = []
            for ci, cell in enumerate(cells):
                if cell is None:
                    continue
                attrs, inner = cell
                explicit = "noref" in attrs or ' r="' in attrs
                ref = "" if explicit else f' r="{col_name(ci)}{rnum}"' if rnum else ""
                attrs = attrs.replace("noref", "")
                cs.append(f"<c{ref}{attrs}>{inner}</c>")
            body.append(f"<row{r_attr}>{''.join(cs)}</row>")
        files[f"xl/worksheets/sheet{i}.xml"] = f'{DECL}<worksheet {NS}><sheetData>{"".join(body)}</sheetData></worksheet>'
        overrides += f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        sheet_entries += f'<sheet name="{escape(sname, {chr(34): "&quot;"})}" sheetId="{i}" r:id="rId{i}"/>'
        wb_rels += f'<Relationship Id="rId{i}" Type="{REL}/worksheet" Target="worksheets/sheet{i}.xml"/>'
    wb_rels += f'<Relationship Id="rIdS" Type="{REL}/styles" Target="styles.xml"/>'
    wb_rels += f'<Relationship Id="rIdT" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/>'
    pr = '<workbookPr date1904="1"/>' if date1904 else "<workbookPr/>"
    files["xl/workbook.xml"] = f"{DECL}<workbook {NS}>{pr}<sheets>{sheet_entries}</sheets></workbook>"
    files["xl/_rels/workbook.xml.rels"] = f'{DECL}<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{wb_rels}</Relationships>'
    nf = "".join(f'<numFmt numFmtId="{i}" formatCode="{escape(c, {chr(34): "&quot;"})}"/>' for i, c in numfmts)
    xf = "".join(f'<xf numFmtId="{i}" fontId="0" fillId="0" borderId="0"/>' for i in xfs)
    files["xl/styles.xml"] = (f'{DECL}<styleSheet {NS}><numFmts count="{len(numfmts)}">{nf}</numFmts>'
                              '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
                              '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
                              '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
                              '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                              f'<cellXfs count="{len(xfs)}">{xf}</cellXfs></styleSheet>')
    files["xl/sharedStrings.xml"] = f'{DECL}<sst {NS}>{"".join(shared)}</sst>'
    files["[Content_Types].xml"] = (
        f'{DECL}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        f"{overrides}</Types>")
    files["_rels/.rels"] = (f'{DECL}<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                            f'<Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    with zipfile.ZipFile(OUT / name, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in files.items():
            z.writestr(zipfile.ZipInfo(n, date_time=(2026, 1, 1, 0, 0, 0)), d)


def n(v, s=None):
    return (f' s="{s}"' if s is not None else "", f"<v>{v}</v>")


def sst(i):
    return (' t="s"', f"<v>{i}</v>")


def inline(text):
    return (' t="inlineStr"', f"<is><t>{escape(text)}</t></is>")


def cell(t, v):
    return (f' t="{t}"', f"<v>{escape(v)}</v>")


def gen_types():
    shared = ["<si><t>Name</t></si>", "<si><t>Value</t></si>", "<si><r><t>rich </t></r><r><rPr><b/></rPr><t>text</t></r></si>",
              "<si><t>esc_x005F_x000D_aped</t></si>", "<si><t xml:space=\"preserve\">  padded  </t></si>",
              "<si><t>NA</t></si>", "<si><t>tab\there</t></si>", '<si><t>a | pipe</t><rPh sb="0" eb="1"><t>phonetic</t></rPh></si>']
    numfmts = [(164, "yyyy-mm-dd"), (165, "[h]:mm:ss"), (166, '"Week "0'), (167, "[Red]0.00"), (168, "d/m/yy\\ h:mm"),
               (169, '0.00_);[Red]\\(0.00\\)'), (170, "mmm d, yyyy")]
    xfs = [0, 14, 164, 165, 166, 167, 22, 21, 27, 168, 169, 170, 10]
    #      0  1   2    3    4    5    6   7   8   9    10   11  12
    rows = [
        (1, [sst(0), sst(1), inline("Kind")]),
        (2, [inline("int"), n(42), inline("plain")]),
        (3, [inline("float"), n(3.14159), inline("plain")]),
        (4, [inline("date14"), n(45292, 1), inline("builtin 14")]),
        (5, [inline("date custom"), n(45292.5, 2), inline("yyyy-mm-dd")]),
        (6, [inline("timedelta"), n(1.5, 3), inline("[h]:mm:ss")]),
        (7, [inline("quoted"), n(12, 4), inline('"Week "0 is not a date')]),
        (8, [inline("colour"), n(1.25, 5), inline("[Red] stripped")]),
        (9, [inline("datetime22"), n(45292.75, 6), inline("m/d/yy h:mm")]),
        (10, [inline("time21"), n(0.5, 7), inline("h:mm:ss")]),
        (11, [inline("cjk27"), n(45292, 8), inline("id 27 is not a date for openpyxl")]),
        (12, [inline("escaped"), n(45292, 9), inline("backslash space")]),
        (13, [inline("bool"), cell("b", "1"), inline("t=b")]),
        (14, [inline("error"), cell("e", "#DIV/0!"), inline("t=e")]),
        (15, [inline("str"), cell("str", "formula text"), inline("t=str")]),
        (16, [inline("isodate"), cell("d", "2024-02-29T13:45:00"), inline("t=d")]),
        (17, [inline("rich"), sst(2), inline("runs joined")]),
        (18, [inline("escape"), sst(3), inline("x005F_ removed")]),
        (19, [inline("padded"), sst(4), inline("strip")]),
        (20, [inline("na string"), sst(5), inline("NA")]),
        (21, [inline("tab"), sst(6), inline("\\t")]),
        (22, [inline("pipe"), sst(7), inline("rPh ignored")]),
        (23, [inline("formula"), (' t="n"', "<f>1+1</f><v>2</v>"), inline("cached")]),
        (24, [inline("empty v"), ("", "<v></v>"), inline("empty")]),
        (25, [inline("percent"), n(0.25, 12), inline("0.00%")]),
        (26, [inline("negative date"), n(-1, 1), inline("before epoch")]),
        (27, [inline("leap1900"), n(59, 1), inline("1900-02-28")]),
        (28, [inline("serial60"), n(60, 1), inline("fake leap day")]),
        (29, [inline("serial61"), n(61, 1), inline("1900-03-01")]),
        (30, [inline("accounting"), n(-3.5, 10), inline("not a date")]),
        (31, [inline("mmm"), n(45300, 11), inline("mmm d, yyyy")]),
    ]
    workbook("gen-cell-types.xlsx", [("Types", rows)], numfmts=numfmts, xfs=xfs, shared=shared)


def gen_columns():
    def col(*vals):
        return vals
    header = ["ints", "floats", "mixed", "strings", "numeric strings", "bools", "bool strings", "dates", "datetimes", "", "dup", "dup", "7"]
    data = [
        [n(1), n(1.5), n(1), inline("a"), inline("007"), cell("b", "1"), inline("TRUE"), n(45292, 1), n(45292.25, 6), n(1), n(1), n(2), n(3)],
        [n(2), n(2.25), inline("x"), None, inline("8"), cell("b", "0"), inline("false"), n(45293, 1), None, None, n(1), n(2), n(3)],
        [n(3), None, n(2.5), inline("c  d"), inline(" 9 "), cell("b", "1"), inline("True"), None, n(45294, 6), n(1), n(1), n(2), n(3)],
        [n(-4), n(1e-07), n(45292, 1), inline("<b>&amp;</b>"), inline("1e3"), None, None, n(45295, 1), n(45295.000011574, 6), None, n(1), n(2), n(3)],
    ]
    rows = [(1, [inline(h) if h else None for h in header])] + [(i + 2, r) for i, r in enumerate(data)]
    workbook("gen-columns.xlsx", [("Columns", rows), ("Empty", []), ("Header only", [(1, [inline("a"), inline("b")])]),
                                  ("Sparse", [(3, [None, None, inline("starts at C3")]), (6, [n(1), None, n(3)]), (7, [])]),
                                  ("One col", [(1, [inline("x")]), (2, [inline("")]), (3, [n(1)]), (4, [inline("  ")]), (5, [n(2)])]),
                                  ("Unicode 表 & <x>", [(1, [inline("名前"), inline("値")]), (2, [inline("東京"), n(1)])])],
             xfs=[0, 14, 164, 165, 166, 167, 22], numfmts=[(164, "yyyy-mm-dd"), (165, "[h]:mm:ss"), (166, "x"), (167, "y")])


def gen_quirks():
    rows = [
        (None, [inline("no r attribute"), inline("on rows")]),
        (None, [inline("second"), n(2)]),
        (5, [("noref", "<v>1</v>"), ("noref", "<v>2</v>"), ("noref", "<v>3</v>")]),
        (4, [inline("out of order row is skipped")]),
        (9, [(' r="D9"', "<v>9</v>"), (' r="B9"', "<v>2</v>")]),
    ]
    workbook("gen-quirks.xlsx", [("Quirks", rows)])
    workbook("gen-date1904.xlsx", [("Mac", [(1, [inline("date"), inline("time")]), (2, [n(0, 1), n(0.75, 2)]), (3, [n(43000, 1), n(1.25, 2)])])],
             xfs=[0, 14, 164], numfmts=[(164, "h:mm")], date1904=True)
    big = [(1, [inline(f"c{i}") for i in range(30)])] + [(r, [n(r * 1000 + i + 0.5) for i in range(30)]) for r in range(2, 60)]
    workbook("gen-wide.xlsx", [("Wide", big)])


def gen_openpyxl():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales 2026"
    ws.append(["Region", "Month", "Units", "Revenue", "Share", "Note"])
    regions = ["North", "South", "East", "West"]
    for i in range(24):
        ws.append([regions[i % 4], dt.date(2026, i % 12 + 1, 1), 100 + i * 7, round(1234.5 + i * 37.25, 2), (i % 5) / 7, None if i % 3 else "check"])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = "yyyy-mm"
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Total units", "=SUM('Sales 2026'!C2:C25)"])
    ws2.append(["Created", dt.datetime(2026, 10, 22, 9, 30)])
    ws2.append(["Ratio", 2 / 3])
    ws2.merge_cells("A6:B6")
    ws2["A6"] = "Merged cell"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["secret", 1])
    wb.save(OUT / "gen-openpyxl.xlsx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.glob("gen-*.xlsx"):
        f.unlink()
    gen_types()
    gen_columns()
    gen_quirks()
    gen_openpyxl()
    print("ok")
