"""Generate the benchmark inputs in bench/corpus/ (deterministic, seed 42).

  medium.docx   50 sections    large.docx  1000 sections
  deck.pptx    300 slides (title, 5-line body, 6x5 table, notes)
  report.pdf   100 pages of Helvetica text (46 lines each, Flate content),
               every fifth page a four-column table

Each docx section = 1 heading + 2 sub-headings + 6 paragraphs with bold and
italic runs + a bullet list + a numbered list + a 6x4 table, mixing English
and Chinese words.

usage: uv run --project oracle python scripts/gen_bench_corpus.py
"""
import random
import zlib
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent.parent / "bench/corpus"
WORDS = ("agent model token markdown parser latency rust python document table "
         "pipeline context retrieval embedding 文档 解析 速度 中文 benchmark").split()


def sentence(rng, n):
    return " ".join(rng.choice(WORDS) for _ in range(n)).capitalize() + "."


def docx(path, sections):
    rng = random.Random(42)
    doc = Document()
    doc.add_heading("Synthetic benchmark document", 0)
    for s in range(sections):
        doc.add_heading(f"Section {s + 1}", 1)
        for sub in range(2):
            doc.add_heading(f"Subsection {s + 1}.{sub + 1}", 2)
            for _ in range(3):
                p = doc.add_paragraph(sentence(rng, 25) + " ")
                p.add_run(sentence(rng, 4)).bold = True
                p.add_run(" " + sentence(rng, 10) + " ")
                p.add_run(sentence(rng, 3)).italic = True
        for _ in range(4):
            doc.add_paragraph(sentence(rng, 8), style="List Bullet")
        for _ in range(3):
            doc.add_paragraph(sentence(rng, 8), style="List Number")
        t = doc.add_table(rows=6, cols=4)
        for r in range(6):
            for c in range(4):
                t.cell(r, c).text = f"h{c}" if r == 0 else f"{rng.choice(WORDS)} {r * c}"
    doc.save(path)


def pptx(path, slides):
    rng = random.Random(42)
    prs = Presentation()
    for i in range(slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide {i} {sentence(rng, 3)}"
        s.placeholders[1].text = "\n".join(sentence(rng, 12) for _ in range(5))
        t = s.shapes.add_table(6, 5, Inches(1), Inches(4), Inches(8), Inches(2)).table
        for r in range(6):
            for c in range(5):
                t.cell(r, c).text = f"{r}-{c} {rng.choice(WORDS)}"
        s.notes_slide.notes_text_frame.text = sentence(rng, 20)
    prs.save(path)


def pdf(path, pages):
    rng = random.Random(42)
    ascii_words = [w for w in WORDS if w.isascii()]
    objs = []
    page_ids = []
    for n in range(pages):
        ops = ["BT /F1 10 Tf 72 740 Td 14 TL"]
        if n % 5 == 4:
            ops.append(f"(Table on page {n + 1}) Tj T*")
            for r in range(30):
                for c in range(4):
                    cell = f"{rng.choice(ascii_words)} {r * c}" if r else f"Column {c}"
                    ops.append(f"q 1 0 0 1 {c * 120} 0 cm ({cell}) Tj Q")
                ops.append("T*")
        else:
            for _ in range(46):
                ops.append(f"({' '.join(rng.choice(ascii_words) for _ in range(12))}) Tj T*")
        ops.append("ET")
        data = zlib.compress("\n".join(ops).encode())
        objs.append(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(data) + data + b"\nendstream")
        content = len(objs) + 3
        objs.append(b"<< /Type /Page /Parent 2 0 R /Contents %d 0 R >>" % content)
        page_ids.append(len(objs) + 3)
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    head = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 3 0 R >> >> >>" % (kids, pages),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.7\n")
    offsets = []
    for i, body in enumerate(head + objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(offsets) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(offsets) + 1, xref)
    path.write_bytes(bytes(out))


OUT.mkdir(parents=True, exist_ok=True)
docx(OUT / "medium.docx", 50)
docx(OUT / "large.docx", 1000)
pptx(OUT / "deck.pptx", 300)
pdf(OUT / "report.pdf", 100)
for p in sorted(OUT.iterdir()):
    print(p.name, p.stat().st_size)
