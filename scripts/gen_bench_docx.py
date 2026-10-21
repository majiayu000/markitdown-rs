"""Generate deterministic synthetic .docx files for benchmarking.

usage: python gen_docx.py <out.docx> <sections>
Each section = 1 heading + 2 sub-headings + 6 paragraphs (bold/italic runs)
+ 1 bullet list + 1 numbered list + 1 table (6x4).
"""
import random
import sys

from docx import Document

WORDS = ("agent model token markdown parser latency rust python document table "
         "pipeline context retrieval embedding 文档 解析 速度 中文 benchmark").split()


def sentence(rng, n):
    return " ".join(rng.choice(WORDS) for _ in range(n)).capitalize() + "."


def main(out, sections):
    rng = random.Random(42)
    doc = Document()
    doc.add_heading("Synthetic benchmark document", 0)
    for s in range(sections):
        doc.add_heading(f"Section {s + 1}", 1)
        for sub in range(2):
            doc.add_heading(f"Subsection {s + 1}.{sub + 1}", 2)
            for _ in range(3):
                p = doc.add_paragraph(sentence(rng, 25) + " ")
                p.add_run(sentence(rng, 4)).bold = True
                p.add_run(" " + sentence(rng, 10) + " ")
                p.add_run(sentence(rng, 3)).italic = True
        for _ in range(4):
            doc.add_paragraph(sentence(rng, 8), style="List Bullet")
        for _ in range(3):
            doc.add_paragraph(sentence(rng, 8), style="List Number")
        t = doc.add_table(rows=6, cols=4)
        for r in range(6):
            for c in range(4):
                t.cell(r, c).text = f"h{c}" if r == 0 else f"{rng.choice(WORDS)} {r * c}"
    doc.save(out)


if __name__ == "__main__":
    main(sys.argv[1], int(sys.argv[2]))
