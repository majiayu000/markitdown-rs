"""Write the generated .pptx fixtures in tests/fixtures/pptx/gen-*.pptx.

Built with python-pptx (deterministic content) and, where python-pptx has no
API for a construct, by editing the slide XML directly.

usage: uv run --project oracle python scripts/gen_pptx_fixtures.py
"""

import copy
import io
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "tests/fixtures/pptx"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
)
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def save(prs, name):
    prs.save(OUT / name)


def text_and_layout():
    prs = Presentation()
    # title + body placeholders inherit their position from the layout
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "   Leading spaces are stripped from titles"
    body = s.placeholders[1].text_frame
    body.text = "First paragraph"
    p = body.add_paragraph()
    p.text = "line one\vline two after a soft break"
    body.add_paragraph().text = ""
    body.add_paragraph().text = "* not a bullet _ not emphasis # not a heading"
    run = body.add_paragraph().add_run()
    run.text = "bold run"
    run.font.bold = True
    # free text boxes are ordered by (top, left), not document order
    for i, (top, left) in enumerate([(5, 5), (1, 6), (5, 1), (1, 1)]):
        tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(1), Inches(0.5))
        tb.text_frame.text = f"box {i} at top={top} left={left}"
    s.notes_slide.notes_text_frame.text = "Speaker notes\nsecond notes line"

    # empty title placeholder, whitespace-only title, notes opened but empty
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "   "
    s.placeholders[1].text = "Body under a blank title"
    s.notes_slide.notes_text_frame.text = "   "

    # title only layout, text in a title-less slide, unicode
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "中文标题 — ünïcode ✓"
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Blank layout, no title placeholder\n\n\nthree newlines above"
    s.shapes.add_connector(1, Inches(0), Inches(0), Inches(2), Inches(2))

    # a placeholder moved on the slide uses its own position
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Moved body"
    s.placeholders[1].text = "body placeholder moved to the top-left"
    s.placeholders[1].top = Emu(0)
    s.placeholders[1].left = Emu(0)
    tb = s.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1))
    tb.text_frame.text = "textbox at the same origin"
    save(prs, "gen-text-layout.pptx")


def groups_and_pictures():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    group = s.shapes.add_group_shape()
    for i, top in enumerate([3, 1, 2]):
        tb = group.shapes.add_textbox(Inches(1), Inches(top), Inches(2), Inches(0.5))
        tb.text_frame.text = f"group child {i} top={top}"
    inner = group.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(4), Inches(0.5), Inches(1), Inches(1)).text_frame.text = "nested group text"
    pic = s.shapes.add_picture(io.BytesIO(PNG), Inches(6), Inches(1), Inches(1), Inches(1))
    pic.name = "Picture 7 (logo) ünï"
    pic._element.nvPicPr.cNvPr.set("descr", "A logo\nwith [brackets]\r\nand   spaces")
    pic2 = s.shapes.add_picture(io.BytesIO(PNG), Inches(6), Inches(3), Inches(1), Inches(1))
    pic2.name = "No description"
    pic3 = group.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(5), Inches(1), Inches(1))
    pic3._element.nvPicPr.cNvPr.set("descr", "   ")

    s = prs.slides.add_slide(prs.slide_layouts[8])  # picture with caption
    s.shapes.title.text = "Picture placeholder"
    s.placeholders[1].insert_picture(io.BytesIO(PNG))
    s.placeholders[2].text = "caption text"
    save(prs, "gen-groups-pictures.pptx")


def tables():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Tables"
    t = s.shapes.add_table(4, 3, Inches(0.5), Inches(1.5), Inches(8), Inches(2)).table
    data = [["Name", "Value | pipe", "Notes"],
            ["a*b_c", "1", "two\nparagraphs"],
            ["", "soft\vbreak", "<b>&amp;</b>"],
            ["merged across", "", "last"]]
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t.cell(3, 0).merge(t.cell(3, 1))
    t2 = s.shapes.add_table(2, 2, Inches(0.5), Inches(4.5), Inches(4), Inches(1)).table
    t2.cell(0, 0).text = "tall"
    t2.cell(0, 0).merge(t2.cell(1, 0))
    t2.cell(0, 1).text = "x"
    t2.cell(1, 1).text = "y"
    save(prs, "gen-tables.pptx")


def charts():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Charts"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Revenue", (1.5, None, 1e-05))
    cd.add_series("Cost", (0.1 + 0.2, 1e16, 3))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(4), Inches(3), cd)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "Clustered | columns"
    xy = XyChartData()
    ser = xy.add_series("Points")
    for x, y in [(1, 2.5), (2, 3.75), (3, -1)]:
        ser.add_data_point(x, y)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(5), Inches(1.5), Inches(4), Inches(3), xy)
    pie = CategoryChartData()
    pie.categories = ["a", "b"]
    pie.add_series("Share", (60, 40))
    gf3 = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(0.5), Inches(5), Inches(3), Inches(2), pie)
    gf3.chart.has_title = False

    # series order by c:order, a missing category and a second plot
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["x", "y", "z"]
    cd.add_series("first", (1, 2, 3))
    cd.add_series("second", (4, 5, 6))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1), Inches(4), Inches(3), cd)
    chart_xml = gf.chart._chartSpace
    c = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    sers = chart_xml.findall(f".//{{{c}}}ser")
    sers[0].find(f"{{{c}}}order").set("val", "5")
    first_cat_pts = sers[0].findall(f"{{{c}}}cat//{{{c}}}pt")
    first_cat_pts[1].getparent().remove(first_cat_pts[1])
    line = chart_xml.find(f".//{{{c}}}lineChart")
    bar = copy.deepcopy(line)
    bar.tag = f"{{{c}}}barChart"
    for ser in bar.findall(f"{{{c}}}ser"):
        ser.find(f"{{{c}}}order").set("val", str(int(ser.find(f"{{{c}}}order").get("val")) + 10))
        for marker in ser.findall(f"{{{c}}}marker") + ser.findall(f"{{{c}}}smooth"):
            ser.remove(marker)
    for el in bar.findall(f"{{{c}}}marker"):
        bar.remove(el)
    bar.insert(0, etree.SubElement(bar, f"{{{c}}}barDir", val="col"))
    line.addnext(bar)

    # unsupported plot types
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["a"]
    cd.add_series("s", (1,))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(3), Inches(3), cd)
    gf.chart._chartSpace.find(f".//{{{c}}}barChart").tag = f"{{{c}}}bar3DChart"
    gf2 = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5), Inches(1), Inches(3), Inches(3), cd)
    plot_area = gf2.chart._chartSpace.find(f".//{{{c}}}plotArea")
    extra = copy.deepcopy(plot_area.find(f"{{{c}}}barChart"))
    extra.tag = f"{{{c}}}bar3DChart"
    plot_area.find(f"{{{c}}}barChart").addnext(extra)
    save(prs, "gen-charts.pptx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.glob("gen-*.pptx"):
        f.unlink()
    text_and_layout()
    groups_and_pictures()
    tables()
    charts()
    for f in sorted(OUT.glob("gen-*.pptx")):
        print(f.name, f.stat().st_size)


if __name__ == "__main__":
    main()
